"""Build the A1 ability model document (business-facing).

    python build_a1_workbook.py [--out a1_ability_reference.xlsx]

The rules follow ``workplace_mirror/specs/A1_ability_spec.md``. Every worked example is run
through the reference rules below and the build FAILS if an example no longer gets the level
or state it illustrates. Results are entry cells (pale yellow): credit cards come from the
check of 31 Aug 2026, retail loans are filled when their check is back. Shares are formulas.
"""
from __future__ import annotations

import argparse
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).parents[1] / "locatability"))
from build_a2_workbook import (  # noqa: E402
    CENTER, INPUT_FILL, MID, MUTED, NAVY, TONE, WINDOW_FILL, WINDOW_HEAD, Sheet, _fill, _fit,
)

# --- rules (reference implementation, monthly granularity) ---------------------------------
W = 3  # months, the scoring month included (61–180 DPD)
OFFSETS = range(1 - 2 * W, 1)  # the prior window and the window, 0 = the scoring month
F_INC, K_INC, F_CASA, K_CASA, F_PAY, K_PAY, DEMIN = 3000, 2, 3000, 1, 500, 1, 1000
FLOORS_IN_TOTAL, D_INC, D_CASA, CASA_MONTHS = 2, 0.25, 0.50, 2


def in_window(k: int) -> bool:
    return k > -W


def income_months(internal: dict, pulls: dict) -> dict:
    """One amount per month: an AECB pull counts in its month and the next, the larger of
    internal and AECB wins."""
    aecb = {**{k + 1: v for k, v in pulls.items()}, **pulls}
    return {k: max(internal.get(k, 0), aecb.get(k, 0)) for k in OFFSETS}


def income_level(internal: dict, pulls: dict, anchor: float = 0) -> str:
    amounts = income_months(internal, pulls)
    now = sum(v for k, v in amounts.items() if in_window(k))
    prior = sum(v for k, v in amounts.items() if not in_window(k))
    if now >= FLOORS_IN_TOTAL * max(F_INC, K_INC * anchor) and now >= prior * (1 - D_INC):
        return "stable"
    return "weak" if any(v >= DEMIN for k, v in amounts.items() if in_window(k)) else "none"


def liquidity_level(casa: dict, payment: float = 0, anchor: float = 0) -> str:
    now = [v for k, v in casa.items() if in_window(k)]
    prior = [v for k, v in casa.items() if not in_window(k)]
    floor = max(F_CASA, K_CASA * anchor)
    holding = not prior or sum(now) / W >= sum(prior) / len(prior) * (1 - D_CASA)
    if (sum(v >= floor for v in now) >= CASA_MONTHS and holding) or payment >= max(F_PAY, K_PAY * anchor):
        return "adequate"
    return "weak" if any(v >= DEMIN for v in now) or payment >= DEMIN else "none"


def ability(income: str, liquidity: str, distress: bool, active: bool) -> str:
    points = {"stable": 2, "weak": 1}.get(income, 0) + {"adequate": 2, "weak": 1}.get(liquidity, 0)
    if points == 0 and not distress and not active:
        return "unknown"
    if points == 4:
        return "moderate" if distress else "high"
    return "moderate" if points >= 2 and not distress else "low"


# --- design ---------------------------------------------------------------------------------
STATES = ["high", "moderate", "low", "unknown"]
LABEL = {"high": "High", "moderate": "Moderate", "low": "Low", "unknown": "Unknown",
         "stable": "Stable", "weak": "Weak", "none": "None",
         "adequate": "Adequate"}
FILLS = {"high": ("A9D08E", "274E13"), "moderate": TONE["good"], "low": TONE["bad"],
         "unknown": TONE["unknown"]}
LEVEL_TONE = {"stable": "high", "adequate": "high", "weak": "moderate", "none": "unknown"}
PULL, CARRIED, CLOSED = ("DDEBF7", "1F4E79"), ("EEF5FB", "1F4E79"), TONE["bad"]


def chip(s: Sheet, col: int, tone: str, label: str | None = None, span: int = 1):
    fill, ink = FILLS[tone]
    return s.text(col, label or LABEL[tone], span=span, bold=True, color=ink, fill=fill, align=MID)


# --- results: (label, tone, credit cards, retail loans), None = not available yet -----------
RESULTS_DATE = "31 Aug 2026"
CC, RL = 19292, None  # accounts
STATE_ROWS = [("high", 47, None), ("moderate", 2689, None), ("low", 5737, None),
              ("unknown", 10819, None)]
RL_PRODUCTS = ["Personal Loans", "Personal Cash Loans", "Auto"]
DETAIL = [
    ("Income: who shows the salary, and the level it reaches", [
        ("Stable · internal salary only", "high", 190, None),
        ("Stable · AECB salary only", "high", 2656, None),
        ("Stable · both", "high", 26, None),
        ("Weak · internal salary only", "moderate", 212, None),
        ("Weak · AECB salary only", "moderate", 1807, None),
        ("Weak · both", "moderate", 12, None),
        ("None · no salary of AED 1,000 or more in the window", "unknown", 14389, None),
    ]),
    ("Income with an AECB salary in the window: why it is Stable or Weak", [
        ("Stable · 1 AECB month in the window", "high", 292, None),
        ("Stable · 2 AECB months", "high", 1225, None),
        ("Stable · 3 AECB months", "high", 1165, None),
        ("Weak · window total reached, but more than 25% below the prior window", "moderate", 912, None),
        ("Weak · window total not reached", "moderate", 907, None),
    ]),
    ("Liquidity: the arm that fired", [
        ("Adequate · CASA balance (no external card payment)", "high", 142, None),
        ("Adequate · with an external card payment (card or CASA arm)", "high", 54, None),
        ("Weak · CASA balance of AED 1,000 or more (no card payment)", "moderate", 327, None),
        ("Weak · with an external card payment", "moderate", 15, None),
        ("None · a card payment below AED 1,000", "unknown", 49, None),
        ("None · no CASA balance or card payment", "unknown", 18705, None),
    ]),
    ("Distress in the window", [
        ("No distress", "moderate", 17616, None),
        ("Bureau bounce only", "low", 483, None),
        ("Cheque return only", "low", 451, None),
        ("Direct-debit return only", "low", 596, None),
        ("Two or more kinds", "low", 146, None),
    ]),
    ("High and Moderate: the combination that gave the state", [
        ("High · Stable income, Adequate liquidity", "high", 47, None),
        ("Moderate · Stable income, no liquidity", "moderate", 2492, None),
        ("Moderate · no income, Adequate liquidity", "moderate", 89, None),
        ("Moderate · Stable income, Weak liquidity", "moderate", 47, None),
        ("Moderate · Weak income, Weak liquidity", "moderate", 28, None),
        ("Moderate · Weak income, Adequate liquidity", "moderate", 20, None),
        ("Moderate · 4 points with distress", "moderate", 13, None),
    ]),
    ("Low: why", [
        ("1 point (mostly Weak income, nothing else)", "low", 2211, None),
        ("0 points, an active CASA where money does not stay", "low", 2065, None),
        ("0 points with distress", "low", 1151, None),
        ("2–3 points with distress", "low", 310, None),
    ]),
    ("Unknown: why", [
        ("Dormant CASA: no customer money movement in the window", "unknown", 6744, None),
        ("No CASA with us (no transaction or positive balance in 12 months)", "unknown", 4075, None),
    ]),
]
SUMMARIES = [
    ("CASA relationship", [
        ("Active: customer moved money in the window", "high", 3631, None),
        ("Dormant: holds a CASA, no customer movement", "unknown", 9125, None),
        ("No CASA with us", "unknown", 6536, None),
    ]),
    ("Internal salary: latest month with AED 1,000 or more", [
        ("In the window", "high", 440, None),
        ("4–6 months ago", "moderate", 3094, None),
        ("7–12 months ago", "moderate", 3599, None),
        ("13–24 months ago", "moderate", 1941, None),
        ("No salary in 24 months (the history available)", "unknown", 10218, None),
    ]),
    ("Latest AECB salary alert of AED 1,000 or more", [
        ("In the window", "high", 3729, None),
        ("4–6 months ago", "moderate", 2362, None),
        ("7–12 months ago", "moderate", 2097, None),
        ("13–24 months ago", "moderate", 1349, None),
        ("25 or more months ago", "moderate", 902, None),
        ("Alert rows, never a salary", "unknown", 7644, None),
        ("No alert row", "unknown", 1209, None),
    ]),
    ("Salary evidence from either source", [
        ("In the window", "high", 4135, None),
        ("Only older", "moderate", 12332, None),
        ("None: no internal salary in 24 months, no AECB salary alert", "unknown", 2825, None),
    ]),
    ("Salary stopped: internal salary in the prior window, none in the window", [
        ("Stopped · no income in the window", "low", 2679, None),
        ("Stopped · Weak on AECB salary", "moderate", 208, None),
        ("Stopped · Stable on AECB salary", "high", 207, None),
        ("Not stopped", "unknown", 16198, None),
    ]),
    ("Highest monthly CASA balance in the window", [
        ("AED 3,000 or more", "high", 242, None),
        ("AED 1,000–3,000", "moderate", 254, None),
        ("Under AED 1,000", "low", 4775, None),
        ("No balance", "unknown", 14021, None),
    ]),
    ("End-of-service result from collections, in the window (does not change the state)", [
        ("Yes", "unknown", 25, None),
    ]),
]
CROSS_COLS = ["AECB in window", "AECB 4–6 m", "AECB 7–12 m", "AECB 13+ m", "Never"]
CROSS = [  # internal salary (rows) x latest AECB salary alert, credit cards
    ("Internal in window", 34, 15, 14, 79, 298),
    ("Internal 4–6 m", 349, 121, 88, 299, 2237),
    ("Internal 7–12 m", 438, 265, 197, 310, 2389),
    ("Internal 13–24 m", 268, 167, 189, 213, 1104),
    ("No internal in 24 m", 2640, 1794, 1609, 1350, 2825),
]
TREND = {  # month: (accounts, high, moderate, low, unknown %)
    "Credit cards 61–180": {
        "Jan": (14664, 0.43, 19.19, 28.48, 51.90), "Feb": (14576, 0.41, 19.40, 27.41, 52.77),
        "Mar": (14455, 0.33, 18.27, 28.19, 53.21), "Apr": (14671, 0.31, 17.84, 28.23, 53.62),
        "May": (15890, 0.35, 17.26, 30.96, 51.44), "Jun": (16982, 0.29, 16.93, 31.39, 51.39),
        "Jul": (18397, 0.29, 15.08, 32.44, 52.19), "Aug": (19292, 0.24, 13.94, 29.74, 56.08),
    },
    "Retail loans 61–180": {m: (None,) * 5 for m in
                            ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug")},
}
NOTES = [
    "AECB salary alerts arrive only when something changes, so no alert means nothing is known, "
    "not that the salary stopped. Pulls every 2 months would give every customer a value and "
    "would show a stop as a pull with no salary.",
    "A customer with fewer alerts in the window than before can look like a fall: 912 accounts "
    "reach the window total but fail the 25% check. Some are real falls; the data cannot yet "
    "split them.",
    "Weak income with nothing else is Low, so a customer with a reported salary can sit below a "
    "customer we know nothing about (Unknown).",
    "The AECB amount is taken as accurate on the pull date. It rarely equals the internal salary "
    "where both exist, and what it looks back over is not yet known.",
    "August is the lowest month for Moderate. The alert data used ends on 27 August, so the month "
    "is slightly short.",
    "Internal salary and balances cover the PRB, PSB and PB segments (95% of base customers). "
    "Internal salary history starts in September 2024.",
]


# --- sheets ---------------------------------------------------------------------------------
def rules(wb: Workbook) -> None:
    s = Sheet(wb.active, [20, 14, 8, 44, 44])
    s.ws.title = "A1 rules"
    s.title("A1 Ability", "State = income points + liquidity points, then distress. Customer-level "
                          "signals are read across all accounts of the customer (CIF).")
    text = ("Internal salary credit, AECB salary, CASA average balance, external card payments, "
            "cheque bounces, DDS (direct debit) returns, bureau bounces, customer-initiated CASA "
            "transactions. Scope: credit cards and retail loans at 61–180 DPD; 180+ DPD follows.")
    s.text(1, "Inputs", bold=True, color=NAVY, border=False)
    s.text(2, text, span=4, border=False)
    _fit(s.ws, s.r, [(text, s.width(2, 5))], floor=18)
    s.r += 1

    s.section("States")
    s.header(["State", "Rule", "In plain words"], [1, 3, 1])
    for st, rule, plain in [
        ("high", "4 points and no distress", "Regular income and a buffer, both confirmed"),
        ("moderate", "≥ 2 but < 4 points and no distress, or 4 points with distress",
         "One side confirmed, or both partly; or everything in place but a payment has just failed"),
        ("low", "< 2 points with any signal, or < 4 points with distress",
         "What we can see shows little or nothing to pay with"),
        ("unknown", "0 points, no distress, and no CASA with us or a dormant one (no "
                    "customer-initiated money movement in the window)",
         "We cannot see the customer's money: absence of data, not evidence of no ability"),
    ]:
        chip(s, 1, st)
        s.text(2, rule, span=3, align=MID)
        s.text(5, plain, align=MID)
        _fit(s.ws, s.r, [(rule, s.width(2, 4)), (plain, s.widths[4])])
        s.r += 1

    s.section("Considerations")
    s.header(["Consideration", "Level", "Points", "Rule"], [1, 1, 1, 2])
    for name, level, points, rule in [
        ("Income (internal salary credit, AECB salary)", "Stable", "2",
         "One salary amount per month, the larger of internal and AECB; an AECB salary counts in its "
         "extract month and the next. Stable = the window's total ≥ 2 × max(AED 3,000, 2 × the amount "
         "due), and the total not more than 25% below the prior window"),
        ("", "Weak", "1", "Some income in the window (AED 1,000 or more in any month) but not Stable: "
                          "below threshold, irregular or weakening"),
        ("", "None", "0", "No internal salary credit and no AECB salary of AED 1,000 or more in any "
                          "month of the window"),
        ("Liquidity (CASA average balance, external card payments)", "Adequate", "2",
         "CASA average balance ≥ max(AED 3,000, 1 × the amount due) in at least 2 of the last 3 months "
         "and the 3-month average not more than 50% below the prior 3 months. Or at least one external "
         "card payment ≥ max(AED 500, 1 × the amount due) in the last 3 months"),
        ("", "Weak", "1", "Not Adequate, but a CASA average balance of AED 1,000 or more in any of the "
                          "last 3 months, or any external card payment of that size"),
        ("", "None", "0", "No CASA average balance and no external card payment of AED 1,000 or more "
                          "in the last 3 months"),
        ("Distress (cheque bounce, DDS return, bureau bounce)", "Yes", "−1 level",
         "1 or more cheque bounces, DDS returns or new bureau bounces in the last 3 months"),
        ("Activity gate (CASA transactions)", "Dormant", "Unknown",
         "Income None, Liquidity None, no distress, and no customer-initiated CASA transaction in the "
         "last 3 months (no CASA at all included)"),
    ]:
        s.text(1, name, bold=True, align=MID)
        s.text(2, level, bold=True, align=MID)
        s.text(3, points, align=CENTER)
        s.text(4, rule, span=2, align=MID)
        _fit(s.ws, s.r, [(name, s.widths[0]), (rule, s.width(4, 5))])
        s.r += 1

    s.section("Map (every combination has one state)")
    s.text(1, "", fill=NAVY)
    s.text(2, "Liquidity Adequate", span=2, bold=True, color="FFFFFF", fill=NAVY, align=CENTER)
    s.text(4, "Liquidity Weak", bold=True, color="FFFFFF", fill=NAVY, align=CENTER)
    s.text(5, "Liquidity None", bold=True, color="FFFFFF", fill=NAVY, align=CENTER)
    s.r += 1
    for distress in (False, True):
        for inc in ("stable", "weak", "none"):
            head = f"{'With' if distress else 'No'} distress · Income {LABEL[inc]}"
            s.text(1, head, bold=True, align=MID)
            for col, span, liq in ((2, 2, "adequate"), (4, 1, "weak"), (5, 1, "none")):
                st = ability(inc, liq, distress, True)
                label = LABEL[st]
                if inc == liq == "none" and not distress:
                    label += " (Unknown if the CASA is dormant or absent)"
                chip(s, col, st, label, span=span).alignment = CENTER
            s.ws.row_dimensions[s.r].height = 30
            s.r += 1

    s.section("Reading the results")
    for n in NOTES:
        s.text(1, n, span=5, border=False)
        _fit(s.ws, s.r, [(n, s.width(1, 5))], floor=18)
        s.r += 1
    s.ws.freeze_panes = "A4"


def thresholds(wb: Workbook) -> None:
    s = Sheet(wb.create_sheet("Thresholds"), [58, 16])
    s.title("Thresholds", "Starting values, not yet tuned on results.")
    s.header(["Threshold", "Value"])
    for name, value in [
        ("Window (the scoring month included), 61–180 DPD / 180+ DPD", "3 / 6 months"),
        ("Months of the window at the CASA floor", "2 of 3"),
        ("Salary floor, internal and AECB", "AED 3,000"),
        ("Income: multiple of the amount due", "2.0"),
        ("Stable: window total, multiples of the income threshold (3 at 180+)", "2"),
        ("Income weakening: fall of the window total vs the prior window", "25%"),
        ("AECB salary: months counted after the extract month", "1"),
        ("CASA average balance floor", "AED 3,000"),
        ("CASA: multiple of the amount due", "1.0"),
        ("CASA weakening: fall of the average vs the prior window", "50%"),
        ("External card payment floor", "AED 500"),
        ("External card payment: multiple of the amount due", "1.0"),
        ("Smallest amount that counts as evidence", "AED 1,000"),
        ("Distress events in the window that count", "1"),
        ("Points needed for High / Moderate", "4 / 2"),
        ("CASA held: months of history for a leg or positive balance", "12"),
    ]:
        s.text(1, name, align=MID)
        s.text(2, value, align=CENTER)
        s.ws.row_dimensions[s.r].height = 20
        s.r += 1
    s.r += 1
    s.note("The amount due: the oldest overdue bucket on cards (none at 180+), one instalment on loans.")
    s.ws.freeze_panes = "A4"


def logic(wb: Workbook) -> None:
    s = Sheet(wb.create_sheet("Logic"), [26, 62, 30])
    s.title("Logic", "How the inputs are read. The window is the scoring month and the 2 before it; "
                     "the prior window is the 3 months before that.")

    def table(title: str, heads: list[str], rows: list, gap: int = 1) -> None:
        s.section(title, gap=gap)
        s.header(heads)
        for row in rows:
            for j, v in enumerate(row, 1):
                s.text(j, v, bold=j == 1, align=MID)
            _fit(s.ws, s.r, [(v, s.widths[j]) for j, v in enumerate(row)])
            s.r += 1

    table("Salary per month", ["Source", "Rule", "Timing"], [
        ("Internal salary", "The month's salary-flagged credits with us, bonuses included.",
         "The month of the credit."),
        ("AECB salary", "The reported average salary on a salary alert (average salary changed, "
                        "salary account added), accurate on the pull date.",
         "The pull's month and the next one, never earlier and no later."),
        ("Salary account closed", "An AECB closure alert ends the AECB salary.", "From the alert."),
        ("One amount per month", "The larger of internal and AECB, never the sum.", "Each month."),
    ], gap=0)
    table("Distress", ["Source", "Rule", "Timing"], [
        ("Cheque / DDS return", "A returned cheque or direct debit on any CASA of the customer.",
         "The month of the return."),
        ("Bureau bounce", "A new bounced amount on a pull; repeated unchanged it counts once.",
         "The pull's month."),
    ])
    table("Customer money movement (the activity gate)", ["", "Transactions", ""], [
        ("Counts", "Transfers in and out, remittances, cash, clearing, bill payments, card spend, "
                   "paying their own credit card.", ""),
        ("Does not count", "The bank recovering the customer's loan instalment, transfers between "
                           "their own accounts, fees, interest, reversals.", ""),
    ])
    s.ws.freeze_panes = "A4"


INCOME_EXAMPLES = [  # level, internal, AECB pulls, anchor, why
    ("stable", {-1: 5000, 0: 5000}, {}, 0, "Two salaries with us, 10,000 against 6,000."),
    ("stable", {-1: 10000}, {}, 0, "Employer paid two salaries in one month: the total counts."),
    ("stable", {-2: 2000, -1: 2000, 0: 2000}, {}, 0, "Three salaries of 2,000 reach the 6,000 total."),
    ("stable", {}, {-2: 3000}, 0, "AECB pull in the window's first month: counts for 2 months."),
    ("stable", {}, {0: 6000}, 0, "AECB pull this month: one month, but 6,000 reaches the total."),
    ("stable", {}, {-2: 3000, 0: 3000}, 0, "Pulls every 2 months fill the window."),
    ("stable", {}, {-3: 6000}, 0, "Pull just before the window carries into its first month."),
    ("weak", {}, {0: 3000}, 0, "AECB pull this month at 3,000: one month, 3,000 against 6,000."),
    ("weak", {-2: 1200, 0: 4000}, {}, 0, "Salary, but 5,200 against the 6,000 total."),
    ("weak", {-1: 5000, 0: 5000}, {}, 3000, "AED 3,000 overdue lifts the total needed to 12,000."),
    ("weak", {-5: 10000, -4: 10000, -3: 10000}, {-1: 10000}, 0,
     "Salary with us stopped; one AECB pull gives 20,000 against 30,000: more than 25% below."),
    ("weak", {}, {-4: 7000, -1: 4000}, 0, "AECB level fell: 8,000 against 14,000."),
    ("weak", {}, {-2: 3000, -1: 0}, 0, "Salary account closed: the AECB amount does not carry."),
    ("none", {0: 800}, {}, 0, "Below AED 1,000."),
    ("none", {}, {-4: 9000}, 0, "AECB pull and its carry are both before the window."),
    ("none", {}, {}, 0, "No salary from either source."),
]
STATE_EXAMPLES = [  # state, internal, pulls, CASA, payment, distress, active, why
    ("high", {-1: 5000, 0: 5000}, {}, {-1: 4000, 0: 4000}, 0, False, True,
     "Salary and balance both confirmed."),
    ("moderate", {}, {-2: 3000}, {}, 0, False, False, "AECB salary, nothing with us."),
    ("moderate", {}, {}, {}, 1500, False, False, "No salary seen, but paying an external card."),
    ("moderate", {-1: 5000, 0: 5000}, {}, {-1: 4000, 0: 4000}, 0, True, True,
     "Everything in place, but a payment just failed."),
    ("low", {}, {0: 3000}, {}, 0, False, False, "Some reported income, nothing else."),
    ("low", {}, {}, {0: 400}, 0, False, True, "Uses the account, but money does not stay."),
    ("low", {-1: 5000, 0: 5000}, {}, {}, 0, True, True, "Stable income, but a cheque bounced."),
    ("low", {}, {}, {}, 0, True, False, "A bureau bounce is all we see."),
    ("unknown", {}, {}, {}, 0, False, False, "No CASA with us, or a dormant one: we cannot see their money."),
]
MONTHS = len(OFFSETS)


def examples(wb: Workbook) -> None:
    s = Sheet(wb.create_sheet("Examples"), [12] + [10] * MONTHS + [62])
    why = MONTHS + 2
    s.title("Examples (61–180 DPD, 3-month window)",
            "One customer per row, oldest month on the left; the blue columns are the window. "
            "Threshold AED 3,000 unless the row says otherwise, so Stable needs a total of 6,000.")
    s.section("Income", gap=0)
    s.text(1, "Cells", bold=True, color=MUTED, border=False, align=MID)
    for col, span, tone, text in ((2, 2, TONE["good"], "5,000  internal salary"),
                                  (4, 2, PULL, "A 3,000  AECB pull"),
                                  (6, 1, CARRIED, "→ carried"), (7, 1, CLOSED, "× closed")):
        s.text(col, text, span=span, color=tone[1], fill=tone[0], align=CENTER)
    s.ws.row_dimensions[s.r].height = 28
    s.r += 1
    s.header(["Income"] + [("Now" if k == 0 else str(k)) for k in OFFSETS] + ["Why"])
    for k in OFFSETS:
        if in_window(k):
            s.ws.cell(row=s.r - 1, column=MONTHS + 1 + k).fill = _fill(WINDOW_HEAD)
    for level, internal, pulls, anchor, text in INCOME_EXAMPLES:
        got = income_level(internal, pulls, anchor)
        if got != level:
            raise SystemExit(f"example '{text}' gives {got}, expected {level}")
        chip(s, 1, LEVEL_TONE[level], LABEL[level])
        for k in OFFSETS:
            col, value, tone = MONTHS + 1 + k, None, None
            if k in pulls:
                value, tone = (f"A {pulls[k]:,}", PULL) if pulls[k] else ("×", CLOSED)
            elif pulls.get(k - 1):
                value, tone = f"→ {pulls[k - 1]:,}", CARRIED
            if internal.get(k, 0) >= (pulls.get(k) or pulls.get(k - 1) or 0) and k in internal:
                value, tone = f"{internal[k]:,}", TONE["good"]
            fill = tone[0] if tone else (WINDOW_FILL if in_window(k) else None)
            s.text(col, value, bold=True, color=tone[1] if tone else "000000", fill=fill, align=CENTER)
        s.text(why, text, align=MID)
        _fit(s.ws, s.r, [(text, s.widths[-1])])
        s.r += 1
    s.section("State")
    s.header(["State", "Income", "Liquidity", "Distress", "Customer money movement", "Why"],
             [1, 2, 1, 1, 2, 1], height=30)
    for state, internal, pulls, casa, payment, distress, active, text in STATE_EXAMPLES:
        inc, liq = income_level(internal, pulls), liquidity_level(casa, payment)
        got = ability(inc, liq, distress, active)
        if got != state:
            raise SystemExit(f"example '{text}' gives {got}, expected {state}")
        chip(s, 1, state)
        chip(s, 2, LEVEL_TONE[inc], LABEL[inc], span=2)
        chip(s, 4, LEVEL_TONE[liq], LABEL[liq])
        s.text(5, "Yes" if distress else "No", align=CENTER)
        s.text(6, "Yes" if active else "No", span=2, align=CENTER)
        s.text(why, text, align=MID)
        _fit(s.ws, s.r, [(text, s.widths[-1])])
        s.r += 1
    s.r += 1
    s.note("At 180+ DPD the window is 6 months and Stable needs 3 × the threshold; the rules are the same.")
    s.ws.freeze_panes = "B4"


def _accounts(s: Sheet, totals: tuple) -> int:
    """An 'Accounts' row, the denominator of the sheet's shares; returns its row."""
    s.text(1, "Accounts", bold=True)
    for j, n in enumerate(totals):
        s.text(2 + 2 * j, n, color="0000FF", fill=INPUT_FILL).number_format = "#,##0"
        s.text(3 + 2 * j, None)
    row = s.r
    s.r += 1
    return row


def _rows(s: Sheet, rows: list, base: int) -> None:
    for label, tone, *counts in rows:
        chip(s, 1, tone, label)
        for j, n in enumerate(counts):
            L = get_column_letter(2 + 2 * j)
            s.text(2 + 2 * j, n, color="0000FF", fill=INPUT_FILL).number_format = "#,##0"
            share = f'=IF(OR({L}{s.r}="",N({L}${base})=0),"",{L}{s.r}/{L}${base})'
            s.text(3 + 2 * j, share, color=MUTED).number_format = "0.0%"
        _fit(s.ws, s.r, [(label, s.widths[0])], floor=18)
        s.r += 1


def _blocks(wb: Workbook, name: str, lead: str, blocks: list) -> Sheet:
    s = Sheet(wb.create_sheet(name), [62, 11, 8, 11, 8])
    s.title(name, lead)
    base = None
    for i, (title, rows) in enumerate(blocks):
        s.section(title, gap=0 if i == 0 else 1)
        s.header(["", "Credit cards", "share", "Retail loans", "share"], height=30)
        if base is None:
            base = _accounts(s, (CC, RL))
        _rows(s, rows, base)
    s.ws.freeze_panes = "A4"
    return s


def results(wb: Workbook) -> None:
    s = Sheet(wb.create_sheet("Results"), [28] + [11, 8] * 5)
    s.title("Results", f"Accounts per state, 61–180 DPD, {RESULTS_DATE}.")
    s.section("Ability state", gap=0)
    heads = ["Credit cards", "Retail loans"] + RL_PRODUCTS
    s.text(1, "State", bold=True, color="FFFFFF", fill=NAVY, align=CENTER)
    for j, h in enumerate(heads):
        s.text(2 + 2 * j, h, span=2, bold=True, color="FFFFFF", fill=NAVY, align=CENTER)
    s.ws.row_dimensions[s.r].height = 30
    s.r += 1
    first, total = s.r, s.r + len(STATE_ROWS)
    for st, cards, loans in STATE_ROWS:
        chip(s, 1, st)
        for j, n in enumerate((cards, loans, None, None, None)):
            L = get_column_letter(2 + 2 * j)
            s.text(2 + 2 * j, n, color="0000FF", fill=INPUT_FILL).number_format = "#,##0"
            share = f'=IF(OR({L}{s.r}="",N({L}${total})=0),"",{L}{s.r}/{L}${total})'
            s.text(3 + 2 * j, share, color=MUTED).number_format = "0.0%"
        s.r += 1
    s.text(1, "Total", bold=True)
    for j in range(len(heads)):
        L = get_column_letter(2 + 2 * j)
        s.text(2 + 2 * j, f'=IF(COUNT({L}{first}:{L}{total - 1})=0,"",SUM({L}{first}:{L}{total - 1}))',
               bold=True).number_format = "#,##0"
        s.text(3 + 2 * j, None)
    s.r += 2
    s.note("Which rule gave each state: 'Results detail'. Coverage of salary, CASA and bureau: "
           "'Results summaries'.")
    s.ws.freeze_panes = "B4"


def detail(wb: Workbook) -> None:
    s = _blocks(wb, "Results detail", f"Which rules fire and why, 61–180 DPD, {RESULTS_DATE}. Shares "
                                      "are of all accounts.", DETAIL)
    s.r += 1
    s.note("The AECB Stable and Weak rows include accounts with internal salary too. A card payment "
           "can sit beside a sufficient balance, so the arms overlap in the second Adequate row.")


def summaries(wb: Workbook) -> None:
    s = _blocks(wb, "Results summaries", f"What we can see of the customer, 61–180 DPD, {RESULTS_DATE}. "
                                         "Shares are of all accounts.", SUMMARIES)
    s.section("Internal salary × AECB salary, credit cards")
    x = Sheet(wb.create_sheet("Salary sources"), [24] + [13] * len(CROSS_COLS))
    wb.move_sheet(x.ws, 0)
    x.title("Salary sources", f"Accounts by the latest internal salary and the latest AECB salary "
                              f"alert, credit cards, {RESULTS_DATE}.")
    x.header([""] + CROSS_COLS, height=30)
    for label, *counts in CROSS:
        x.text(1, label, bold=True, align=MID)
        for j, n in enumerate(counts, 2):
            x.text(j, n, color="0000FF", fill=INPUT_FILL).number_format = "#,##0"
        x.r += 1
    x.r += 1
    x.note("'Never' = no AECB salary alert of AED 1,000 or more, with or without alert rows. "
           "Retail loans follow.")
    s.text(1, "See the sheet 'Salary sources'.", span=5, border=False)
    s.r += 1


def trend(wb: Workbook) -> None:
    s = Sheet(wb.create_sheet("Results trend"), [12, 12] + [11] * len(STATES))
    s.title("Results trend", "Share of accounts per state by month, Jan–Aug 2026 (month ends), "
                             "61–180 DPD.")
    for i, (title, months) in enumerate(TREND.items()):
        s.section(title, gap=0 if i == 0 else 1)
        s.text(1, "Month", bold=True, color="FFFFFF", fill=NAVY, align=CENTER)
        s.text(2, "Accounts", bold=True, color="FFFFFF", fill=NAVY, align=CENTER)
        for j, st in enumerate(STATES):
            chip(s, 3 + j, st).alignment = CENTER
        s.ws.row_dimensions[s.r].height = 20
        s.r += 1
        for month, (n, *shares) in months.items():
            s.text(1, f"{month} 2026", bold=True)
            s.text(2, n, color="0000FF", fill=INPUT_FILL).number_format = "#,##0"
            for j, v in enumerate(shares):
                cell = s.text(3 + j, None if v is None else v / 100, color="0000FF", fill=INPUT_FILL)
                cell.number_format = "0.0%"
            s.r += 1
    s.r += 1
    s.note("The book grows from May. Moderate falls steadily, and August also has slightly short "
           "alert data ('A1 rules').")
    s.ws.freeze_panes = "A4"


def build(out: pathlib.Path) -> pathlib.Path:
    wb = Workbook()
    rules(wb)
    thresholds(wb)
    logic(wb)
    examples(wb)
    results(wb)
    detail(wb)
    summaries(wb)
    trend(wb)
    wb.move_sheet("Salary sources", offset=len(wb.sheetnames) - 2 - wb.sheetnames.index("Salary sources"))
    wb.calculation.fullCalcOnLoad = True
    for ws in wb:
        ws.sheet_properties.tabColor = NAVY if ws.title in ("A1 rules", "Thresholds", "Logic", "Examples") else "2E75B6"
    wb.save(out)
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", default=str(pathlib.Path(__file__).with_name("a1_ability_reference.xlsx")))
    print(build(pathlib.Path(ap.parse_args().out)))


if __name__ == "__main__":
    main()
